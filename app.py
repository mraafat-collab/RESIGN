"""
Resignation Approval Portal  (v2 – rebuilt)
-------------------------------------------
Streamlit + SQLite.  Beige / brown design.  Template e-mails from the company mailbox.
Run:  streamlit run app.py
"""
import base64
import hmac
import io
import os
import re
import shutil
import smtplib
import sqlite3
import ssl
import tempfile
from contextlib import closing
from datetime import date, datetime, timedelta
from email.mime.application import MIMEApplication
from email.mime.image import MIMEImage
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formataddr
from html import escape
from pathlib import Path

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from docx import Document

# =========================================================
# Config
# =========================================================
st.set_page_config(
    page_title="Resignation Portal",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded",
)

APP_DIR = Path(__file__).parent
DB_PATH = Path(os.environ.get("DB_PATH", APP_DIR / "portal_data.db"))

PALETTES = {
    "Espresso": dict(dark="#3E2710", brown="#6B4423", mid="#8C5A2B", accent="#C08A4B",
                     bg="#F7F2EA", card="#FFFDF9", border="#D6C2A7", soft="#EADBCE",
                     text="#3B2A1A", muted="#8a7560"),
    "Mocha": dict(dark="#4A3426", brown="#7A5238", mid="#A06E4A", accent="#D2A679",
                  bg="#F6EFE8", card="#FFFCF8", border="#DCC7B3", soft="#EFE1D4",
                  text="#3F2D22", muted="#907A68"),
    "Sand": dict(dark="#5C4631", brown="#8A6A47", mid="#B08B5E", accent="#D9B77E",
                 bg="#FAF6EE", card="#FFFEFA", border="#E0CFB2", soft="#F1E6D2",
                 text="#45362A", muted="#9A8566"),
}
P = PALETTES["Espresso"]

STATUS_AR = {"Pending": "قيد المراجعة", "Approved": "معتمدة", "Rejected": "مرفوضة"}
STATUS_ICON = {"Pending": "⏳", "Approved": "✅", "Rejected": "⛔"}
HIRE_STATUS_AR = {"Planned": "مخطط", "Onboarding": "تحت التجهيز", "Active": "على رأس العمل"}
DEFAULT_FOOTER = "Best regards,\nHR Department"

RES_FIELDS = [
    "employee_code", "employee_name", "job_title", "branch", "department", "grading",
    "hiring_date", "resignation_date", "last_working_day", "net_pay", "reason",
    "direct_manager", "recipient_name", "notice_period", "handover_status", "notes",
]
HIRE_FIELDS = [
    "full_name", "job_title", "branch", "department", "start_date", "salary",
    "manager", "phone", "email", "status", "replaces_id", "notes",
]


# =========================================================
# Secrets
# =========================================================
def get_secret(key, default=""):
    try:
        if key in st.secrets:
            return st.secrets[key]
    except Exception:
        pass
    return os.environ.get(key, default)


SENDER_EMAIL = get_secret("SENDER_EMAIL")
SENDER_PASSWORD = get_secret("SENDER_PASSWORD")
SENDER_NAME = get_secret("SENDER_NAME", "HR Department")
DEFAULT_SMTP = get_secret("SMTP_SERVER", "smtp.office365.com")
DEFAULT_PORT = int(get_secret("SMTP_PORT", 587))
APP_PASSWORD = get_secret("APP_PASSWORD")


# =========================================================
# Database
# =========================================================
def _conn():
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c


def run(sql, params=()):
    with closing(_conn()) as c:
        cur = c.execute(sql, params)
        c.commit()
        return cur.lastrowid


def fetch(sql, params=()):
    with closing(_conn()) as c:
        return [dict(r) for r in c.execute(sql, params).fetchall()]


def init_db():
    with closing(_conn()) as c:
        c.executescript(
            """
            CREATE TABLE IF NOT EXISTS resignations(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                employee_code TEXT, employee_name TEXT NOT NULL, job_title TEXT, branch TEXT,
                department TEXT, grading TEXT, hiring_date TEXT, resignation_date TEXT,
                last_working_day TEXT, net_pay REAL DEFAULT 0, reason TEXT, direct_manager TEXT,
                recipient_name TEXT, notice_period TEXT, handover_status TEXT,
                status TEXT DEFAULT 'Pending', notes TEXT, photo BLOB,
                created_at TEXT, updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS hires(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                full_name TEXT NOT NULL, job_title TEXT, branch TEXT, department TEXT,
                start_date TEXT, salary REAL DEFAULT 0, manager TEXT, phone TEXT, email TEXT,
                status TEXT DEFAULT 'Planned', replaces_id INTEGER, notes TEXT,
                created_at TEXT, updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS audit_log(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ts TEXT, action TEXT, entity TEXT, entity_id INTEGER, details TEXT
            );
            CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
            """
        )
        c.commit()


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def log(action, entity, entity_id, details=""):
    run(
        "INSERT INTO audit_log(ts, action, entity, entity_id, details) VALUES(?,?,?,?,?)",
        (now_iso(), action, entity, entity_id, details),
    )


def get_setting(key, default=""):
    rows = fetch("SELECT value FROM settings WHERE key=?", (key,))
    return rows[0]["value"] if rows else default


def set_setting(key, value):
    run(
        "INSERT INTO settings(key, value) VALUES(?,?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, value),
    )


def to_date(s):
    try:
        return date.fromisoformat(str(s)[:10])
    except Exception:
        return date.today()


def load_resignations():
    rows = fetch("SELECT * FROM resignations ORDER BY id DESC")
    for r in rows:
        for k in ("hiring_date", "resignation_date", "last_working_day"):
            r[k] = to_date(r[k])
        r["net_pay"] = int(r["net_pay"] or 0)
        r["exp_days"] = max((r["last_working_day"] - r["hiring_date"]).days, 0)
        for k in ("employee_code", "job_title", "branch", "department", "grading", "reason",
                  "direct_manager", "recipient_name", "notice_period", "handover_status", "notes"):
            r[k] = r[k] or ""
    return rows


def load_hires():
    rows = fetch("SELECT * FROM hires ORDER BY id DESC")
    for r in rows:
        r["start_date"] = to_date(r["start_date"])
        r["salary"] = int(r["salary"] or 0)
        for k in ("job_title", "branch", "department", "manager", "phone", "email", "notes"):
            r[k] = r[k] or ""
    return rows


def save_resignation(vals, rid=None, photo=None):
    vals = dict(vals)
    for k in ("hiring_date", "resignation_date", "last_working_day"):
        vals[k] = vals[k].isoformat()
    cols = list(RES_FIELDS)
    params = [vals[k] for k in cols]
    if photo is not None:
        cols.append("photo")
        params.append(photo)
    if rid:
        sets = ", ".join(f"{c}=?" for c in cols) + ", updated_at=?"
        run(f"UPDATE resignations SET {sets} WHERE id=?", params + [now_iso(), rid])
        log("تعديل", "resignation", rid, vals["employee_name"])
        return rid
    cols += ["status", "created_at", "updated_at"]
    params += ["Pending", now_iso(), now_iso()]
    ph = ",".join("?" * len(cols))
    new_id = run(f"INSERT INTO resignations({','.join(cols)}) VALUES({ph})", params)
    log("إضافة", "resignation", new_id, vals["employee_name"])
    return new_id


def save_hire(vals, hid=None):
    vals = dict(vals)
    vals["start_date"] = vals["start_date"].isoformat()
    cols = list(HIRE_FIELDS)
    params = [vals[k] for k in cols]
    if hid:
        sets = ", ".join(f"{c}=?" for c in cols) + ", updated_at=?"
        run(f"UPDATE hires SET {sets} WHERE id=?", params + [now_iso(), hid])
        log("تعديل", "hire", hid, vals["full_name"])
        return hid
    cols += ["created_at", "updated_at"]
    params += [now_iso(), now_iso()]
    ph = ",".join("?" * len(cols))
    new_id = run(f"INSERT INTO hires({','.join(cols)}) VALUES({ph})", params)
    log("إضافة", "hire", new_id, vals["full_name"])
    return new_id


def restore_db(raw):
    if not raw.startswith(b"SQLite format 3\x00"):
        return False, "الملف ليس قاعدة بيانات SQLite صالحة."
    with tempfile.NamedTemporaryFile(delete=False, suffix=".db") as tmp:
        tmp.write(raw)
        tmp_path = tmp.name
    try:
        with closing(sqlite3.connect(tmp_path)) as c:
            tables = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        need = {"resignations", "hires", "audit_log", "settings"}
        if not need.issubset(tables):
            return False, "هذه ليست نسخة احتياطية من هذا النظام (جداول ناقصة)."
        if DB_PATH.exists():
            shutil.copy(DB_PATH, str(DB_PATH) + ".bak")
        shutil.copy(tmp_path, DB_PATH)
        return True, "تمت الاستعادة بنجاح."
    finally:
        os.unlink(tmp_path)


def seed_demo():
    today = date.today()
    demo = [
        ("Abu Bakr Muhammad Mahmoud", "Sales Casual", "Park-St Branch", "Education", 7000, 25, 40, "Approved"),
        ("Mona Hassan Ali", "Cashier", "City Center", "Relocation", 5200, 12, 20, "Pending"),
        ("Omar Khaled Samir", "Team Leader", "Park-St Branch", "Better Offer", 11800, 5, 14, "Pending"),
        ("Salma Adel Fathy", "Kitchen Staff", "Nasr City", "Personal", 4300, 60, 30, "Rejected"),
    ]
    for i, (name, title, br, reason, pay, ago, notice, status) in enumerate(demo, start=1):
        rid = save_resignation(
            dict(
                employee_code=f"E-{1000 + i}", employee_name=name, job_title=title, branch=br,
                department="Operations", grading="G1 - L1",
                hiring_date=today - timedelta(days=500 + i * 90),
                resignation_date=today - timedelta(days=ago),
                last_working_day=today - timedelta(days=ago) + timedelta(days=notice),
                net_pay=pay, reason=reason, direct_manager="Sherif S. El-Sayed",
                recipient_name="Mohamed", notice_period="One Month", handover_status="Team Leader",
                notes="",
            )
        )
        if status != "Pending":
            run("UPDATE resignations SET status=? WHERE id=?", (status, rid))
    save_hire(
        dict(full_name="Youssef Tarek", job_title="Sales Casual", branch="Park-St Branch",
             department="Operations", start_date=today + timedelta(days=10), salary=6500,
             manager="Sherif S. El-Sayed", phone="", email="", status="Planned", replaces_id=1, notes="")
    )
    save_hire(
        dict(full_name="Nour Ibrahim", job_title="Cashier", branch="City Center",
             department="Operations", start_date=today - timedelta(days=7), salary=5400,
             manager="", phone="", email="", status="Active", replaces_id=None, notes="")
    )
    log("بيانات تجريبية", "system", 0, "تم تحميل بيانات تجريبية")


# =========================================================
# Small helpers
# =========================================================
def to_english_str(text):
    if not text:
        return ""
    return str(text).translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")).strip()


def fmt(d):
    return d.strftime("%d/%m/%Y")


def split_emails(value):
    return [e.strip() for e in re.split(r"[;,\s]+", value or "") if e.strip()]


def valid_email(addr):
    return re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", addr) is not None


def recipient_title_of(r):
    rec = (r.get("recipient_name") or "Manager").strip() or "Manager"
    return rec if rec.lower().startswith(("mr.", "ms.", "mrs.", "dr.")) else f"Mr. {rec}"


def tenure_text(days):
    y, rem = divmod(days, 365)
    m, d = divmod(rem, 30)
    parts = []
    if y:
        parts.append(f"{y} سنة")
    if m:
        parts.append(f"{m} شهر")
    if d or not parts:
        parts.append(f"{d} يوم")
    return " و ".join(parts)


def badge(status):
    colors = {
        "Pending": ("#FFF3D6", "#8A5A00", "#F0D9A0"),
        "Approved": ("#E7F0DC", "#44631F", "#C6D9A8"),
        "Rejected": ("#F7E0DA", "#8A2E1B", "#E8B8AC"),
    }
    bg, fg, bd = colors.get(status, ("#EEE", "#333", "#CCC"))
    return (f'<span style="background:{bg};color:{fg};border:1px solid {bd};padding:3px 12px;'
            f'border-radius:999px;font-size:12px;font-weight:700;">'
            f'{STATUS_ICON.get(status, "")} {STATUS_AR.get(status, status)}</span>')


def kpi(label, value, icon="", size=26):
    return (f'<div class="kpi"><div class="l">{icon} {escape(str(label))}</div>'
            f'<div class="v" style="font-size:{size}px">{escape(str(value))}</div></div>')


def show_df(df, **kw):
    try:
        st.dataframe(df, width="stretch", hide_index=True, **kw)
    except Exception:
        st.dataframe(df, use_container_width=True, hide_index=True, **kw)


def flash(msg):
    st.session_state["_flash"] = msg


def img_to_b64(raw):
    return base64.b64encode(raw).decode() if raw else None


# =========================================================
# Exports
# =========================================================
def generate_word_doc(r):
    doc = Document()
    doc.add_heading(f"RESIGNATION APPROVAL REQUEST - {r['employee_name']}", 0)
    rows = [
        ("Recipient", f"Dear {recipient_title_of(r)}"),
        ("Employee Name", r["employee_name"]), ("Employee Code", r["employee_code"]),
        ("Job Title", r["job_title"]), ("Branch", r["branch"]), ("Department", r["department"]),
        ("Hiring Date", fmt(r["hiring_date"])), ("Resignation Date", fmt(r["resignation_date"])),
        ("Last Working Day", fmt(r["last_working_day"])),
        ("Total Experience", f"{r['exp_days']} Days"),
        ("Total Net Pay", f"{r['net_pay']:,} EGP"), ("Reason for Leaving", r["reason"]),
        ("Direct Manager", r["direct_manager"]), ("Handover Status", r["handover_status"]),
        ("Status", r["status"]),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, str(v)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_excel(rs):
    df = pd.DataFrame(
        [
            {
                "ID": r["id"], "Employee Code": r["employee_code"], "Employee Name": r["employee_name"],
                "Job Title": r["job_title"], "Branch": r["branch"], "Department": r["department"],
                "Hiring Date": fmt(r["hiring_date"]), "Resignation Date": fmt(r["resignation_date"]),
                "Last Working Day": fmt(r["last_working_day"]), "Experience (Days)": r["exp_days"],
                "Net Pay (EGP)": r["net_pay"], "Reason": r["reason"],
                "Direct Manager": r["direct_manager"], "Recipient": r["recipient_name"],
                "Status": r["status"],
            }
            for r in rs
        ]
    )
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="xlsxwriter") as writer:
        df.to_excel(writer, index=False, sheet_name="Resignations")
        ws = writer.sheets["Resignations"]
        head = writer.book.add_format(
            {"bold": True, "bg_color": P["mid"], "font_color": "#FFFFFF", "border": 1}
        )
        for i, col in enumerate(df.columns):
            ws.write(0, i, col, head)
            ws.set_column(i, i, max(16, len(col) + 4))
        ws.freeze_panes(1, 0)
    return out.getvalue()


# =========================================================
# E-mail template
# =========================================================
def build_email_html(r, note="", photo_src=None, kind="request"):
    p = P
    e = lambda v: escape(str(v if v is not None else ""))
    company = get_setting("company_name", "Company")
    footer = get_setting("email_footer", DEFAULT_FOOTER)
    status = r["status"]

    if kind == "request":
        title = "RESIGNATION APPROVAL REQUEST"
        intro = "Kindly review the resignation request below and confirm your approval."
    else:
        title = f"RESIGNATION {status.upper()}"
        intro = {
            "Approved": "Please be informed that the resignation request below has been approved.",
            "Rejected": "Please be informed that the resignation request below has been rejected.",
            "Pending": "Please be informed that the resignation request below is still under review.",
        }[status]

    photo_cell = (
        f'<img src="{photo_src}" width="88" height="114" '
        f'style="border-radius:8px;border:1px solid {p["border"]};object-fit:cover;display:block;" alt="photo">'
        if photo_src
        else f'<div style="width:88px;height:114px;background:{p["soft"]};border-radius:8px;'
             f'border:1px solid {p["border"]};text-align:center;line-height:114px;font-size:34px;">&#128100;</div>'
    )

    def row(label, value, shade):
        bg = "#FAF8F5" if shade else "#FFFFFF"
        return (
            f'<tr><td style="padding:9px 14px;background:{bg};border-bottom:1px solid {p["soft"]};'
            f'font-size:13px;color:#7a6650;width:38%;">{label}</td>'
            f'<td style="padding:9px 14px;background:{bg};border-bottom:1px solid {p["soft"]};'
            f'font-size:13px;color:#222;font-weight:600;">{value}</td></tr>'
        )

    rows = [
        ("Employee Code", e(r["employee_code"])), ("Employee Name", e(r["employee_name"])),
        ("Job Title", e(r["job_title"])), ("Branch", e(r["branch"])),
        ("Hiring Date", fmt(r["hiring_date"])), ("Resignation Date", fmt(r["resignation_date"])),
        ("Last Working Day", fmt(r["last_working_day"])), ("Total Experience", f"{r['exp_days']} Days"),
        ("Reason for Leaving", e(r["reason"])), ("Direct Manager", e(r["direct_manager"])),
        ("Handover Status", e(r["handover_status"])),
    ]
    rows = [x for x in rows if x[1].strip()]
    rows_html = "".join(row(k, v, i % 2 == 0) for i, (k, v) in enumerate(rows))

    note_html = (
        f'<tr><td style="padding:0 28px 18px;"><div style="background:#FFF6E8;border-left:4px solid {p["accent"]};'
        f'padding:12px 16px;border-radius:8px;font-size:13px;color:{p["text"]};line-height:1.6;">'
        f'{e(note).replace(chr(10), "<br>")}</div></td></tr>'
        if note.strip() else ""
    )
    footer_html = e(footer).replace(chr(10), "<br>")

    return f"""<!DOCTYPE html>
<html><body style="margin:0;padding:24px 12px;background:{p['bg']};font-family:Arial,Helvetica,sans-serif;">
<table role="presentation" width="640" align="center" cellpadding="0" cellspacing="0"
 style="max-width:640px;width:100%;background:{p['card']};border:1px solid {p['border']};border-radius:14px;overflow:hidden;">
  <tr><td style="background:{p['mid']};padding:22px 28px;">
    <div style="color:#F3E3CC;font-size:11px;letter-spacing:2px;text-transform:uppercase;">{e(company)}</div>
    <div style="color:#F3E3CC;font-size:13px;margin-top:8px;">Dear {e(recipient_title_of(r))},</div>
    <div style="color:#FFFFFF;font-size:22px;font-weight:800;letter-spacing:.5px;margin-top:4px;">{title}</div>
  </td></tr>
  <tr><td style="padding:18px 28px 6px;font-size:13px;color:{p['text']};line-height:1.6;">{intro}</td></tr>
  <tr><td style="padding:12px 28px;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0"><tr>
      <td width="100" valign="top">{photo_cell}</td>
      <td valign="middle" style="padding-left:16px;">
        <div style="font-size:18px;font-weight:800;color:#111;">{e(r['employee_name'])}</div>
        <div style="font-size:13px;color:#555;margin-top:2px;">{e(r['job_title'])}</div>
        <div style="font-size:12px;color:#888;">{e(r['branch'])}</div>
      </td></tr></table>
  </td></tr>
  <tr><td style="padding:0 28px 18px;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="6"><tr>
      <td align="center" style="background:{p['mid']};color:#fff;border-radius:10px;padding:10px 6px;">
        <div style="font-size:10px;opacity:.85;">NET PAY</div>
        <div style="font-size:15px;font-weight:800;">{r['net_pay']:,} EGP</div></td>
      <td align="center" style="background:{p['dark']};color:#fff;border-radius:10px;padding:10px 6px;">
        <div style="font-size:10px;opacity:.85;">LAST WORK DAY</div>
        <div style="font-size:15px;font-weight:800;">{fmt(r['last_working_day'])}</div></td>
      <td align="center" style="background:{p['accent']};color:#fff;border-radius:10px;padding:10px 6px;">
        <div style="font-size:10px;opacity:.9;">EXPERIENCE</div>
        <div style="font-size:15px;font-weight:800;">{r['exp_days']} Days</div></td>
    </tr></table>
  </td></tr>
  <tr><td style="padding:0 28px 18px;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0"
     style="border:1px solid {p['border']};border-radius:10px;overflow:hidden;">{rows_html}</table>
  </td></tr>
  {note_html}
  <tr><td style="padding:4px 28px 18px;font-size:13px;color:{p['text']};line-height:1.6;">{footer_html}</td></tr>
  <tr><td style="padding:12px 28px;background:{p['soft']};font-size:11px;color:#7a6650;text-align:center;">
    {e(company)} &bull; Resignation Portal &bull; {date.today().strftime('%d/%m/%Y')}
  </td></tr>
</table></body></html>"""


def build_plain_text(r, note="", kind="request"):
    head = "RESIGNATION APPROVAL REQUEST" if kind == "request" else f"RESIGNATION {r['status'].upper()}"
    lines = [
        f"Dear {recipient_title_of(r)},", "", head, "",
        f"Employee: {r['employee_name']}", f"Job Title: {r['job_title']}", f"Branch: {r['branch']}",
        f"Hiring Date: {fmt(r['hiring_date'])}", f"Resignation Date: {fmt(r['resignation_date'])}",
        f"Last Working Day: {fmt(r['last_working_day'])}", f"Total Experience: {r['exp_days']} Days",
        f"Total Net Pay: {r['net_pay']:,} EGP", f"Reason: {r['reason']}",
        f"Direct Manager: {r['direct_manager']}",
    ]
    if note.strip():
        lines += ["", note.strip()]
    lines += ["", get_setting("email_footer", DEFAULT_FOOTER)]
    return "\n".join(lines)


def connect_smtp(server, port):
    if port == 465:
        s = smtplib.SMTP_SSL(server, port, timeout=20, context=ssl.create_default_context())
    else:
        s = smtplib.SMTP(server, port, timeout=20)
        s.ehlo()
        s.starttls(context=ssl.create_default_context())
        s.ehlo()
    s.login(SENDER_EMAIL, SENDER_PASSWORD)
    return s


def send_email(r, to_list, cc_list, subject, note, kind, server, port, attach_word, attach_excel):
    photo = r.get("photo")
    html = build_email_html(r, note, "cid:employee_photo" if photo else None, kind)

    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"] = formataddr((SENDER_NAME, SENDER_EMAIL))
    msg["To"] = ", ".join(to_list)
    if cc_list:
        msg["Cc"] = ", ".join(cc_list)

    related, alt = MIMEMultipart("related"), MIMEMultipart("alternative")
    alt.attach(MIMEText(build_plain_text(r, note, kind), "plain", "utf-8"))
    alt.attach(MIMEText(html, "html", "utf-8"))
    related.attach(alt)
    if photo:
        img = MIMEImage(photo)
        img.add_header("Content-ID", "<employee_photo>")
        img.add_header("Content-Disposition", "inline", filename="photo")
        related.attach(img)
    msg.attach(related)

    safe = re.sub(r"[^\w\-]+", "_", r["employee_name"]) or "employee"
    if attach_word:
        part = MIMEApplication(generate_word_doc(r), Name=f"Resignation_{safe}.docx")
        part.add_header("Content-Disposition", "attachment", filename=f"Resignation_{safe}.docx")
        msg.attach(part)
    if attach_excel:
        part = MIMEApplication(generate_excel([r]), Name=f"Resignation_{safe}.xlsx")
        part.add_header("Content-Disposition", "attachment", filename=f"Resignation_{safe}.xlsx")
        msg.attach(part)

    s = connect_smtp(server, port)
    try:
        s.sendmail(SENDER_EMAIL, to_list + cc_list, msg.as_string())
    finally:
        s.quit()


# =========================================================
# Boot: DB + theme + CSS
# =========================================================
init_db()
theme_name = get_setting("theme", "Espresso")
if theme_name not in PALETTES:
    theme_name = "Espresso"
P = PALETTES[theme_name]

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;600;700;800&display=swap');
html, body, [class*="css"], .stApp { font-family: 'Cairo', sans-serif; color: @text@; }
.stApp {
  background:
    radial-gradient(900px 500px at 100% -10%, rgba(192,138,75,.16), transparent 60%),
    radial-gradient(700px 400px at -10% 110%, rgba(140,90,43,.10), transparent 60%), @bg@;
}
#MainMenu, footer { visibility: hidden; }
header[data-testid="stHeader"] { background: transparent; }
.block-container { padding-top: 1.2rem; max-width: 1450px; }

/* sidebar */
section[data-testid="stSidebar"] { background: @card@; border-right: 1px solid @border@; }
.brand { background: linear-gradient(135deg, @dark@, @mid@); color: #FFF8EE; border-radius: 18px;
  padding: 16px 18px; margin-bottom: 14px; direction: rtl; text-align: right; }
.brand b { font-size: 18px; } .brand span { display:block; font-size: 12px; opacity: .8; }

/* arabic labels */
[data-testid="stWidgetLabel"], [data-testid="stWidgetLabel"] p { direction: rtl; text-align: right; width: 100%; font-weight: 600; }
.stMarkdown h1, .stMarkdown h2, .stMarkdown h3, .stMarkdown h4, .stMarkdown p { text-align: start; unicode-bidi: plaintext; }

/* hero */
.hero { background: linear-gradient(120deg, @dark@ 0%, @brown@ 45%, @accent@ 130%);
  border-radius: 26px; padding: 26px 34px; margin-bottom: 14px; color: #FFF8EE; position: relative;
  overflow: hidden; box-shadow: 0 16px 40px rgba(62,39,16,.28); direction: rtl; text-align: right; }
.hero:after { content:""; position:absolute; left:-70px; top:-70px; width:240px; height:240px;
  border-radius:50%; background: rgba(255,255,255,.07); }
.hero .eyebrow { font-size: 12px; letter-spacing: 3px; text-transform: uppercase; opacity: .75; }
.hero h1 { margin: 4px 0 0; font-size: 30px; font-weight: 800; color: #FFF8EE; }
.hero p { margin: 6px 0 0; opacity: .85; font-size: 14px; }

/* cards */
.card { background: @card@; border: 1px solid @border@; border-radius: 20px; padding: 18px 22px;
  margin-bottom: 14px; box-shadow: 0 6px 22px rgba(140,90,43,.09); direction: rtl; text-align: right; }
.card h4 { margin: 0 0 6px; color: @dark@; font-weight: 800; }
.muted { color: @muted@; font-size: 13px; }
.kpi { background: @card@; border: 1px solid @border@; border-radius: 18px; padding: 14px 18px;
  text-align: right; direction: rtl; border-right: 6px solid @mid@; box-shadow: 0 4px 16px rgba(140,90,43,.08); }
.kpi .l { font-size: 12px; color: @muted@; font-weight: 700; }
.kpi .v { font-weight: 800; color: @brown@; margin-top: 2px; }

/* detail grid */
.dgrid { display: grid; grid-template-columns: repeat(3, 1fr); gap: 10px 18px; direction: rtl; text-align: right; }
.dgrid div small { display:block; color: @muted@; font-size: 11px; font-weight: 700; }
.dgrid div span { font-weight: 700; color: @text@; unicode-bidi: plaintext; }

/* timeline */
.tl { display:flex; align-items:flex-start; gap:0; margin: 8px 0 4px; direction: ltr; }
.tl .st { flex:1; text-align:center; position:relative; }
.tl .st:not(:last-child):after { content:""; position:absolute; top:13px; left:50%; width:100%; height:3px; background:@border@; z-index:0; }
.tl .dot { width:28px; height:28px; border-radius:50%; background:@mid@; color:#fff; margin:0 auto; display:flex;
  align-items:center; justify-content:center; font-size:13px; position:relative; z-index:1; border:3px solid @card@; }
.tl .st b { display:block; font-size:12px; margin-top:6px; color:@dark@; }
.tl .st small { color:@muted@; font-size:11px; }

/* tabs */
.stTabs [data-baseweb="tab-list"] { gap: 4px; background: @card@; padding: 6px; border-radius: 18px;
  border: 1px solid @border@; flex-wrap: wrap; box-shadow: 0 4px 16px rgba(140,90,43,.07); }
.stTabs [data-baseweb="tab"] { height: 46px; border-radius: 12px; padding: 0 16px; color: @dark@; background: transparent; }
.stTabs [data-baseweb="tab"] p { font-size: 15px; font-weight: 700; margin: 0; color: inherit; }
.stTabs [data-baseweb="tab"]:hover { background: @soft@; }
.stTabs [aria-selected="true"] { background: linear-gradient(135deg, @mid@, @dark@) !important; color: #FFF8EE !important;
  box-shadow: 0 6px 16px rgba(62,39,16,.30); }
.stTabs [aria-selected="true"] p { color: #FFF8EE !important; }
.stTabs [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-border"] { display: none; }

/* buttons */
.stButton > button, .stDownloadButton > button, .stFormSubmitButton > button {
  background: linear-gradient(135deg, @mid@, @dark@); color: #FFF8EE; font-weight: 700; border: none;
  border-radius: 12px; padding: 9px 20px; box-shadow: 0 6px 16px rgba(62,39,16,.22); transition: all .22s ease; width: 100%; }
.stButton > button:hover, .stDownloadButton > button:hover, .stFormSubmitButton > button:hover {
  transform: translateY(-2px); color: #fff; background: linear-gradient(135deg, @accent@, @mid@); }
.stButton > button p, .stDownloadButton > button p, .stFormSubmitButton > button p { color: #FFF8EE; }

/* inputs */
.stTextInput input, .stTextArea textarea, .stNumberInput input, .stDateInput input, div[data-baseweb="select"] > div {
  background: #fff !important; border-color: @border@ !important; border-radius: 10px !important; color: @text@; }
div[data-baseweb="popover"] li { color: @text@; }
.stTextInput input:focus, .stTextArea textarea:focus { border-color: @mid@ !important; box-shadow: 0 0 0 2px rgba(140,90,43,.18) !important; }
details[data-testid="stExpander"] { background: @card@; border: 1px solid @border@; border-radius: 16px; }
details[data-testid="stExpander"] summary p { font-weight: 700; color: @dark@; }
.pill { display:inline-block; padding:4px 12px; border-radius:999px; font-size:12px; font-weight:700; }
.pill.ok { background:#E7F0DC; color:#44631F; border:1px solid #C6D9A8; }
.pill.bad { background:#F7E0DA; color:#8A2E1B; border:1px solid #E8B8AC; }
</style>
"""
for _k, _v in P.items():
    CSS = CSS.replace(f"@{_k}@", _v)
st.markdown(CSS, unsafe_allow_html=True)

# =========================================================
# Optional password gate
# =========================================================
if APP_PASSWORD and not st.session_state.get("auth_ok"):
    _, mid, _ = st.columns([1, 1, 1])
    with mid:
        st.markdown('<div class="card"><h4>🔐 تسجيل الدخول</h4><div class="muted">أدخل كلمة المرور للوصول للنظام</div></div>',
                    unsafe_allow_html=True)
        pw = st.text_input("كلمة المرور", type="password", key="login_pw")
        if st.button("دخول", key="login_btn"):
            if hmac.compare_digest(pw.encode(), str(APP_PASSWORD).encode()):
                st.session_state["auth_ok"] = True
                st.rerun()
            else:
                st.error("كلمة المرور غير صحيحة.")
    st.stop()

# =========================================================
# Sidebar: settings, backup
# =========================================================
company_name = get_setting("company_name", "Company")
with st.sidebar:
    st.markdown(
        f'<div class="brand"><b>🏛️ {escape(company_name)}</b><span>بوابة الاستقالات والتعيينات</span></div>',
        unsafe_allow_html=True,
    )

    st.markdown("##### ⚙️ الإعدادات")
    new_company = st.text_input("اسم الشركة", value=company_name, key="set_company")
    new_theme = st.selectbox("🎨 اللون", list(PALETTES), index=list(PALETTES).index(theme_name), key="set_theme")
    new_footer = st.text_area("توقيع الرسائل", value=get_setting("email_footer", DEFAULT_FOOTER),
                              height=90, key="set_footer")
    if st.button("💾 حفظ الإعدادات", key="save_settings"):
        set_setting("company_name", new_company.strip() or "Company")
        set_setting("theme", new_theme)
        set_setting("email_footer", new_footer)
        log("إعدادات", "system", 0, "تم تحديث الإعدادات")
        st.rerun()

    st.markdown("---")
    st.markdown("##### ✉️ حالة البريد")
    if SENDER_EMAIL and SENDER_PASSWORD:
        st.markdown('<span class="pill ok">● مُهيّأ</span>', unsafe_allow_html=True)
        st.caption(SENDER_EMAIL)
    else:
        st.markdown('<span class="pill bad">● غير مُهيّأ</span>', unsafe_allow_html=True)
        st.caption("أضف SENDER_EMAIL و SENDER_PASSWORD في Secrets")

    st.markdown("---")
    st.markdown("##### 🗄️ النسخ الاحتياطي")
    st.caption("على Streamlit Cloud الملفات مؤقتة. نزّل نسخة بعد كل تعديل مهم.")
    try:
        st.download_button("⬇️ تنزيل نسخة احتياطية", DB_PATH.read_bytes(),
                           file_name=f"portal_backup_{date.today()}.db",
                           mime="application/octet-stream", key="dl_backup")
    except Exception:
        pass
    up = st.file_uploader("استعادة من نسخة (.db)", type=["db"], key="restore_up")
    if up is not None and st.button("♻️ استعادة الآن", key="restore_btn"):
        ok, msg = restore_db(up.getvalue())
        if ok:
            flash(msg)
            st.rerun()
        else:
            st.error(msg)

# =========================================================
# Data + hero
# =========================================================
recs = load_resignations()
hires = load_hires()
rec_map = {r["id"]: r for r in recs}

n_total = len(recs)
n_pending = sum(r["status"] == "Pending" for r in recs)
n_appr = sum(r["status"] == "Approved" for r in recs)
n_rej = sum(r["status"] == "Rejected" for r in recs)

st.markdown(
    f"""
<div class="hero">
  <div class="eyebrow">{escape(company_name)}</div>
  <h1>Resignation Approval System</h1>
  <p>تابع الاستقالات، اعتمدها، أرسل المراسلات الرسمية، وأرشف كل شيء في مكان واحد.</p>
</div>""",
    unsafe_allow_html=True,
)

if "_flash" in st.session_state:
    st.success(st.session_state.pop("_flash"))

# Global current record
current = None
if recs:
    sel_col, badge_col = st.columns([4, 1])
    with sel_col:
        cur_id = st.selectbox(
            "🎯 السجل الحالي (يُستخدم في: الاستقالات، ملف الموظف، البريد)",
            list(rec_map),
            format_func=lambda i: f"{rec_map[i]['employee_name']}  ·  {rec_map[i]['branch']}  ·  {STATUS_AR[rec_map[i]['status']]}",
            key="current_record",
        )
    current = rec_map[cur_id]
    with badge_col:
        st.markdown(f'<div style="padding-top:30px;text-align:center">{badge(current["status"])}</div>',
                    unsafe_allow_html=True)

tab_dash, tab_res, tab_hires, tab_file, tab_mail, tab_an, tab_log = st.tabs(
    ["🏠 لوحة القيادة", "📄 الاستقالات", "🧑‍💼 التعيينات الجديدة", "📁 ملف الموظف",
     "✉️ البريد الرسمي", "📊 التحليلات", "🧾 سجل النشاط"]
)


# =========================================================
# Forms
# =========================================================
def resignation_form(key, rec=None):
    r = rec or {}
    d0 = date.today()
    lo, hi = date(1990, 1, 1), date(2100, 12, 31)
    with st.form(f"res_form_{key}", clear_on_submit=(rec is None)):
        c1, c2, c3 = st.columns(3)
        name = c1.text_input("اسم الموظف *", value=r.get("employee_name", ""))
        code = c1.text_input("الكود الوظيفي", value=r.get("employee_code", ""))
        title = c1.text_input("المسمى الوظيفي", value=r.get("job_title", ""))
        branch = c2.text_input("الفرع", value=r.get("branch", ""))
        dept = c2.text_input("القسم", value=r.get("department", ""))
        grading = c2.text_input("الدرجة (Grading)", value=r.get("grading", ""))
        mgr = c3.text_input("المدير المباشر", value=r.get("direct_manager", ""))
        recip = c3.text_input("الموجه له الطلب (Recipient)", value=r.get("recipient_name", ""))
        pay = c3.number_input("إجمالي المستحقات (EGP)", min_value=0, value=int(r.get("net_pay", 0)), step=100)
        d1, d2, d3 = st.columns(3)
        hd = d1.date_input("تاريخ التعيين", value=r.get("hiring_date", d0), min_value=lo, max_value=hi)
        rd = d2.date_input("تاريخ الاستقالة", value=r.get("resignation_date", d0), min_value=lo, max_value=hi)
        ld = d3.date_input("آخر يوم عمل", value=r.get("last_working_day", d0), min_value=lo, max_value=hi)
        e1, e2, e3 = st.columns(3)
        reason = e1.text_input("سبب الاستقالة", value=r.get("reason", ""))
        notice = e2.text_input("فترة الإشعار", value=r.get("notice_period", "") or "One Month")
        handover = e3.text_input("حالة التسليم (Handover)", value=r.get("handover_status", ""))
        notes = st.text_area("ملاحظات داخلية", value=r.get("notes", ""), height=70)
        photo = st.file_uploader("الصورة الشخصية (اختياري)", type=["png", "jpg", "jpeg"], key=f"photo_{key}")
        submitted = st.form_submit_button("💾 حفظ")
    vals = dict(
        employee_code=to_english_str(code), employee_name=to_english_str(name),
        job_title=to_english_str(title), branch=to_english_str(branch),
        department=to_english_str(dept), grading=to_english_str(grading),
        hiring_date=hd, resignation_date=rd, last_working_day=ld, net_pay=int(pay),
        reason=to_english_str(reason), direct_manager=to_english_str(mgr),
        recipient_name=to_english_str(recip), notice_period=to_english_str(notice),
        handover_status=to_english_str(handover), notes=notes.strip(),
    )
    return submitted, vals, (photo.getvalue() if photo else None)


def validate_resignation(vals, photo):
    if not vals["employee_name"]:
        return "اسم الموظف مطلوب."
    if vals["resignation_date"] < vals["hiring_date"]:
        return "تاريخ الاستقالة لا يمكن أن يسبق تاريخ التعيين."
    if vals["last_working_day"] < vals["resignation_date"]:
        return "آخر يوم عمل لا يمكن أن يسبق تاريخ الاستقالة."
    if photo is not None and len(photo) > 3 * 1024 * 1024:
        return "الصورة أكبر من 3MB."
    return None


def hire_form(key, rec=None):
    r = rec or {}
    lo, hi = date(1990, 1, 1), date(2100, 12, 31)
    options = [None] + list(rec_map)
    cur_rep = r.get("replaces_id")
    idx = options.index(cur_rep) if cur_rep in options else 0
    with st.form(f"hire_form_{key}", clear_on_submit=(rec is None)):
        c1, c2, c3 = st.columns(3)
        name = c1.text_input("الاسم الكامل *", value=r.get("full_name", ""))
        title = c1.text_input("المسمى الوظيفي", value=r.get("job_title", ""))
        phone = c1.text_input("الموبايل", value=r.get("phone", ""))
        branch = c2.text_input("الفرع", value=r.get("branch", ""))
        dept = c2.text_input("القسم", value=r.get("department", ""))
        mail = c2.text_input("الإيميل", value=r.get("email", ""))
        mgr = c3.text_input("المدير المباشر", value=r.get("manager", ""))
        start = c3.date_input("تاريخ البدء", value=r.get("start_date", date.today()), min_value=lo, max_value=hi)
        salary = c3.number_input("الراتب (EGP)", min_value=0, value=int(r.get("salary", 0)), step=100)
        f1, f2 = st.columns(2)
        status = f1.selectbox("الحالة", list(HIRE_STATUS_AR), index=list(HIRE_STATUS_AR).index(r.get("status", "Planned")),
                              format_func=lambda s: HIRE_STATUS_AR[s])
        repl = f2.selectbox("يحل محل (اختياري)", options, index=idx,
                            format_func=lambda i: "—" if i is None else f"{rec_map[i]['employee_name']} ({rec_map[i]['branch']})")
        notes = st.text_area("ملاحظات", value=r.get("notes", ""), height=60)
        submitted = st.form_submit_button("💾 حفظ")
    vals = dict(
        full_name=to_english_str(name), job_title=to_english_str(title), branch=to_english_str(branch),
        department=to_english_str(dept), start_date=start, salary=int(salary),
        manager=to_english_str(mgr), phone=to_english_str(phone), email=mail.strip(),
        status=status, replaces_id=repl, notes=notes.strip(),
    )
    return submitted, vals


# =========================================================
# TAB: Dashboard
# =========================================================
with tab_dash:
    if not recs and not hires:
        st.markdown('<div class="card"><h4>👋 أهلاً بك</h4><div class="muted">لا توجد بيانات بعد. أضف أول استقالة من تبويب «الاستقالات»، أو جرّب النظام ببيانات تجريبية.</div></div>',
                    unsafe_allow_html=True)
        if st.button("🧪 تحميل بيانات تجريبية", key="seed_demo"):
            seed_demo()
            flash("تم تحميل بيانات تجريبية.")
            st.rerun()
    else:
        k = st.columns(5)
        k[0].markdown(kpi("إجمالي الاستقالات", n_total, "📄"), unsafe_allow_html=True)
        k[1].markdown(kpi("قيد المراجعة", n_pending, "⏳"), unsafe_allow_html=True)
        k[2].markdown(kpi("معتمدة", n_appr, "✅"), unsafe_allow_html=True)
        k[3].markdown(kpi("مرفوضة", n_rej, "⛔"), unsafe_allow_html=True)
        k[4].markdown(kpi("تعيينات جديدة", len(hires), "🧑‍💼"), unsafe_allow_html=True)
        st.write("")

        left, right = st.columns([1.6, 1])
        with left:
            st.markdown('<div class="card"><h4>🕑 آخر الطلبات</h4></div>', unsafe_allow_html=True)
            if recs:
                show_df(pd.DataFrame([
                    {"الموظف": r["employee_name"], "الفرع": r["branch"], "آخر يوم": r["last_working_day"],
                     "المستحقات": r["net_pay"], "الحالة": STATUS_AR[r["status"]]}
                    for r in recs[:8]
                ]))
        with right:
            st.markdown('<div class="card"><h4>🔔 تنبيهات</h4></div>', unsafe_allow_html=True)
            soon = [r for r in recs if r["status"] != "Rejected"
                    and date.today() <= r["last_working_day"] <= date.today() + timedelta(days=14)]
            if n_pending:
                st.warning(f"{n_pending} طلب بانتظار القرار.")
            if soon:
                for r in soon:
                    days_left = (r["last_working_day"] - date.today()).days
                    st.info(f"**{r['employee_name']}** — آخر يوم بعد {days_left} يوم ({fmt(r['last_working_day'])})")
            if not n_pending and not soon:
                st.success("لا توجد تنبيهات عاجلة.")

        if recs:
            g1, g2 = st.columns(2)
            df_d = pd.DataFrame([{"branch": r["branch"] or "—", "reason": r["reason"] or "—"} for r in recs])
            with g1:
                st.markdown('<div class="card"><h4>الاستقالات حسب الفرع</h4></div>', unsafe_allow_html=True)
                st.bar_chart(df_d["branch"].value_counts(), color=P["mid"])
            with g2:
                st.markdown('<div class="card"><h4>الاستقالات حسب السبب</h4></div>', unsafe_allow_html=True)
                st.bar_chart(df_d["reason"].value_counts(), color=P["accent"])

# =========================================================
# TAB: Resignations
# =========================================================
with tab_res:
    # ---- list ----
    st.markdown('<div class="card"><h4>📄 كل الاستقالات</h4></div>', unsafe_allow_html=True)
    if recs:
        f1, f2, f3 = st.columns([2, 1.3, 1.3])
        q = f1.text_input("🔎 بحث بالاسم / الكود / الوظيفة", key="f_q")
        f_status = f2.multiselect("الحالة", list(STATUS_AR), format_func=lambda s: STATUS_AR[s], key="f_status")
        branches = sorted({r["branch"] for r in recs if r["branch"]})
        f_branch = f3.multiselect("الفرع", branches, key="f_branch")

        shown = [
            r for r in recs
            if (not q or q.lower() in f"{r['employee_name']} {r['employee_code']} {r['job_title']}".lower())
            and (not f_status or r["status"] in f_status)
            and (not f_branch or r["branch"] in f_branch)
        ]
        st.caption(f"عرض {len(shown)} من {len(recs)}")
        if shown:
            show_df(pd.DataFrame([
                {"#": r["id"], "الموظف": r["employee_name"], "الكود": r["employee_code"],
                 "الوظيفة": r["job_title"], "الفرع": r["branch"],
                 "تاريخ الاستقالة": r["resignation_date"], "آخر يوم": r["last_working_day"],
                 "المستحقات": r["net_pay"], "الحالة": STATUS_AR[r["status"]], "السبب": r["reason"]}
                for r in shown
            ]))
            st.download_button("📊 تصدير القائمة (Excel)", generate_excel(shown),
                               file_name=f"resignations_{date.today()}.xlsx", key="dl_all_xlsx")
    else:
        st.info("لا توجد استقالات بعد.")

    # ---- action panel for current ----
    if current:
        st.markdown("---")
        st.markdown(
            f'<div class="card"><h4>🎛️ إجراءات على: {escape(current["employee_name"])} {badge(current["status"])}</h4>'
            f'<div class="muted">{escape(current["job_title"])} · {escape(current["branch"])} · '
            f'آخر يوم عمل {fmt(current["last_working_day"])} · مستحقات {current["net_pay"]:,} EGP</div></div>',
            unsafe_allow_html=True,
        )
        a1, a2, a3 = st.columns(3)
        for col, (st_key, label) in zip((a1, a2, a3), [("Approved", "✅ اعتماد"), ("Rejected", "⛔ رفض"), ("Pending", "↩️ إرجاع للمراجعة")]):
            if col.button(label, key=f"status_{st_key}", disabled=(current["status"] == st_key)):
                run("UPDATE resignations SET status=?, updated_at=? WHERE id=?", (st_key, now_iso(), current["id"]))
                log("تغيير الحالة", "resignation", current["id"], f"{current['employee_name']} → {STATUS_AR[st_key]}")
                flash(f"تم تغيير حالة {current['employee_name']} إلى {STATUS_AR[st_key]}.")
                st.rerun()

        with st.expander("✏️ تعديل بيانات هذا السجل"):
            ok, vals, photo = resignation_form(f"edit_{current['id']}", current)
            if ok:
                err = validate_resignation(vals, photo)
                if err:
                    st.error(err)
                else:
                    save_resignation(vals, current["id"], photo)
                    flash("تم حفظ التعديلات.")
                    st.rerun()

        with st.expander("🗑️ حذف هذا السجل"):
            st.warning("الحذف نهائي ولا يمكن التراجع عنه (إلا باستعادة نسخة احتياطية).")
            sure = st.checkbox("أؤكد رغبتي في الحذف", key=f"del_ok_{current['id']}")
            if st.button("❌ حذف نهائي", key=f"del_btn_{current['id']}", disabled=not sure):
                run("DELETE FROM resignations WHERE id=?", (current["id"],))
                run("UPDATE hires SET replaces_id=NULL WHERE replaces_id=?", (current["id"],))
                log("حذف", "resignation", current["id"], current["employee_name"])
                flash("تم حذف السجل.")
                st.rerun()

    # ---- add ----
    st.markdown("---")
    with st.expander("➕ إضافة استقالة جديدة", expanded=not recs):
        ok, vals, photo = resignation_form("add")
        if ok:
            err = validate_resignation(vals, photo)
            if err:
                st.error(err)
            else:
                save_resignation(vals, None, photo)
                flash("تمت إضافة الاستقالة بحالة «قيد المراجعة».")
                st.rerun()

# =========================================================
# TAB: Hires
# =========================================================
with tab_hires:
    h = st.columns(3)
    h[0].markdown(kpi("مخطط", sum(x["status"] == "Planned" for x in hires), "🗓️"), unsafe_allow_html=True)
    h[1].markdown(kpi("تحت التجهيز", sum(x["status"] == "Onboarding" for x in hires), "🧰"), unsafe_allow_html=True)
    h[2].markdown(kpi("على رأس العمل", sum(x["status"] == "Active" for x in hires), "🟢"), unsafe_allow_html=True)
    st.write("")

    st.markdown('<div class="card"><h4>🧑‍💼 قائمة التعيينات</h4></div>', unsafe_allow_html=True)
    if hires:
        show_df(pd.DataFrame([
            {"#": x["id"], "الاسم": x["full_name"], "الوظيفة": x["job_title"], "الفرع": x["branch"],
             "تاريخ البدء": x["start_date"], "الراتب": x["salary"], "الحالة": HIRE_STATUS_AR[x["status"]],
             "يحل محل": rec_map[x["replaces_id"]]["employee_name"] if x["replaces_id"] in rec_map else "—"}
            for x in hires
        ]))
        hmap = {x["id"]: x for x in hires}
        hid = st.selectbox("اختر تعيين للتعديل أو الحذف", list(hmap), key="hire_pick",
                           format_func=lambda i: f"{hmap[i]['full_name']} · {HIRE_STATUS_AR[hmap[i]['status']]}")
        with st.expander("✏️ تعديل التعيين المحدد"):
            ok, vals = hire_form(f"edit_{hid}", hmap[hid])
            if ok:
                if not vals["full_name"]:
                    st.error("الاسم مطلوب.")
                else:
                    save_hire(vals, hid)
                    flash("تم حفظ التعديلات.")
                    st.rerun()
        with st.expander("🗑️ حذف التعيين المحدد"):
            sure_h = st.checkbox("أؤكد الحذف", key=f"hdel_ok_{hid}")
            if st.button("❌ حذف", key=f"hdel_btn_{hid}", disabled=not sure_h):
                run("DELETE FROM hires WHERE id=?", (hid,))
                log("حذف", "hire", hid, hmap[hid]["full_name"])
                flash("تم حذف التعيين.")
                st.rerun()
    else:
        st.info("لا توجد تعيينات بعد.")

    with st.expander("➕ إضافة تعيين جديد", expanded=not hires):
        ok, vals = hire_form("add")
        if ok:
            if not vals["full_name"]:
                st.error("الاسم مطلوب.")
            else:
                save_hire(vals, None)
                flash("تمت إضافة التعيين.")
                st.rerun()

# =========================================================
# TAB: Employee file
# =========================================================
with tab_file:
    if not current:
        st.info("اختر سجلاً أولاً.")
    else:
        r = current
        pc, ic = st.columns([1, 4])
        with pc:
            if r["photo"]:
                st.image(r["photo"], width=140)
            else:
                st.markdown(f'<div style="width:140px;height:180px;border-radius:14px;background:{P["soft"]};'
                            f'border:1px solid {P["border"]};display:flex;align-items:center;justify-content:center;'
                            f'font-size:56px;">👤</div>', unsafe_allow_html=True)
        with ic:
            st.markdown(
                f"""<div class="card">
<h4 style="font-size:22px">{escape(r['employee_name'])} {badge(r['status'])}</h4>
<div class="muted" style="margin-bottom:12px">{escape(r['job_title'])} · {escape(r['branch'])} · {escape(r['department'])}</div>
<div class="dgrid">
 <div><small>الكود الوظيفي</small><span>{escape(r['employee_code']) or '—'}</span></div>
 <div><small>الدرجة</small><span>{escape(r['grading']) or '—'}</span></div>
 <div><small>المدير المباشر</small><span>{escape(r['direct_manager']) or '—'}</span></div>
 <div><small>مدة الخدمة</small><span>{tenure_text(r['exp_days'])}</span></div>
 <div><small>إجمالي المستحقات</small><span>{r['net_pay']:,} EGP</span></div>
 <div><small>سبب الاستقالة</small><span>{escape(r['reason']) or '—'}</span></div>
 <div><small>فترة الإشعار</small><span>{escape(r['notice_period']) or '—'}</span></div>
 <div><small>حالة التسليم</small><span>{escape(r['handover_status']) or '—'}</span></div>
 <div><small>الموجه له الطلب</small><span>{escape(r['recipient_name']) or '—'}</span></div>
</div></div>""",
                unsafe_allow_html=True,
            )

        st.markdown(
            f"""<div class="card"><h4>🧭 المسار الزمني</h4>
<div class="tl">
 <div class="st"><div class="dot">1</div><b>التعيين</b><small>{fmt(r['hiring_date'])}</small></div>
 <div class="st"><div class="dot">2</div><b>تقديم الاستقالة</b><small>{fmt(r['resignation_date'])}</small></div>
 <div class="st"><div class="dot">3</div><b>آخر يوم عمل</b><small>{fmt(r['last_working_day'])}</small></div>
</div></div>""",
            unsafe_allow_html=True,
        )

        if r["notes"]:
            st.markdown(f'<div class="card"><h4>📝 ملاحظات</h4><div>{escape(r["notes"])}</div></div>', unsafe_allow_html=True)

        repl = [x for x in hires if x["replaces_id"] == r["id"]]
        if repl:
            st.markdown('<div class="card"><h4>🔁 البديل المعيّن</h4>'
                        + "".join(f'<div>{escape(x["full_name"])} — {HIRE_STATUS_AR[x["status"]]} — يبدأ {fmt(x["start_date"])}</div>' for x in repl)
                        + "</div>", unsafe_allow_html=True)

        safe = re.sub(r"[^\w\-]+", "_", r["employee_name"])
        d1, d2, _ = st.columns([1, 1, 2])
        d1.download_button("📄 تحميل Word", generate_word_doc(r), file_name=f"Resignation_{safe}.docx", key="dl_word")
        d2.download_button("📊 تحميل Excel", generate_excel([r]), file_name=f"Resignation_{safe}.xlsx", key="dl_xlsx")

        with st.expander("🖨️ كارت A4 للطباعة / تنزيل صورة PNG", expanded=False):
            src = f"data:image/png;base64,{img_to_b64(r['photo'])}" if r["photo"] else None
            card = build_email_html(r, "", src, "request")
            page = f"""
            <script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script>
            <style>
              body{{margin:0;background:transparent;font-family:Arial,sans-serif}}
              .bar{{display:flex;gap:8px;justify-content:flex-end;margin-bottom:10px}}
              .b{{background:linear-gradient(135deg,{P['mid']},{P['dark']});color:#fff;border:0;padding:9px 18px;
                  border-radius:10px;font-weight:700;cursor:pointer}}
              @media print{{.bar{{display:none}}}}
            </style>
            <div class="bar">
              <button class="b" onclick="dl()">🖼️ تنزيل PNG</button>
              <button class="b" onclick="window.print()">🖨️ طباعة</button>
            </div>
            <div id="cap">{card}</div>
            <script>
              function dl(){{
                html2canvas(document.getElementById('cap'),{{scale:2,useCORS:true}}).then(c=>{{
                  const a=document.createElement('a');a.download="Resignation_{safe}.png";
                  a.href=c.toDataURL('image/png');a.click();}});
              }}
            </script>"""
            components.html(page, height=820, scrolling=True)

        hist = fetch("SELECT ts, action, details FROM audit_log WHERE entity='resignation' AND entity_id=? ORDER BY id DESC LIMIT 50",
                     (r["id"],))
        st.markdown('<div class="card"><h4>🧾 سجل هذا الموظف</h4></div>', unsafe_allow_html=True)
        if hist:
            show_df(pd.DataFrame(hist).rename(columns={"ts": "الوقت", "action": "الإجراء", "details": "التفاصيل"}))
        else:
            st.caption("لا يوجد نشاط مسجل.")

# =========================================================
# TAB: Mail
# =========================================================
with tab_mail:
    if not current:
        st.info("اختر سجلاً أولاً.")
    else:
        r = current
        left, right = st.columns([1, 1.5])
        with left:
            st.markdown('<div class="card"><h4>✉️ إعداد الرسالة</h4>'
                        '<div class="muted">تُرسل من إيميل الشركة المُعرّف في Secrets</div></div>',
                        unsafe_allow_html=True)
            kind_label = st.radio("نوع القالب", ["طلب موافقة على استقالة", "إشعار بقرار الاستقالة"], key="mail_kind")
            kind = "request" if kind_label.startswith("طلب") else "decision"
            default_subject = (f"Resignation Approval Request - {r['employee_name']}" if kind == "request"
                               else f"Resignation {r['status']} - {r['employee_name']}")
            to_raw = st.text_input("إلى (أكثر من إيميل بفاصلة)", key="mail_to", placeholder="hr@company.com, boss@company.com")
            cc_raw = st.text_input("نسخة (CC)", key="mail_cc")
            subject = st.text_input("الموضوع", value=default_subject, key=f"mail_subj_{kind}_{r['id']}")
            note = st.text_area("ملاحظة إضافية داخل الرسالة (اختياري)", height=90, key="mail_note")
            m1, m2 = st.columns(2)
            att_w = m1.checkbox("إرفاق Word", value=True, key="att_w")
            att_x = m2.checkbox("إرفاق Excel", value=False, key="att_x")

            with st.expander("⚙️ إعدادات السيرفر"):
                smtp_server = st.text_input("SMTP Server", value=DEFAULT_SMTP, key="smtp_srv")
                ports = [587, 465, 25]
                smtp_port = st.selectbox("Port", ports, index=ports.index(DEFAULT_PORT) if DEFAULT_PORT in ports else 0, key="smtp_port")
                if st.button("🔌 اختبار الاتصال", key="smtp_test"):
                    if not (SENDER_EMAIL and SENDER_PASSWORD):
                        st.error("بيانات الإيميل غير موجودة في Secrets.")
                    else:
                        try:
                            with st.spinner("جاري الاتصال..."):
                                s = connect_smtp(smtp_server.strip(), int(smtp_port))
                                s.quit()
                            st.success("✅ الاتصال وتسجيل الدخول ناجحان")
                        except Exception as ex:
                            st.error(f"❌ {ex}")

            if st.button("🚀 إرسال الآن", key="mail_send"):
                to_list, cc_list = split_emails(to_raw), split_emails(cc_raw)
                bad = [x for x in to_list + cc_list if not valid_email(x)]
                if not (SENDER_EMAIL and SENDER_PASSWORD):
                    st.error("أضف SENDER_EMAIL و SENDER_PASSWORD في Secrets.")
                elif not to_list:
                    st.error("اكتب إيميل المستلم.")
                elif bad:
                    st.error(f"إيميل غير صحيح: {', '.join(bad)}")
                else:
                    try:
                        with st.spinner("جاري الإرسال..."):
                            send_email(r, to_list, cc_list, subject, note, kind,
                                       smtp_server.strip(), int(smtp_port), att_w, att_x)
                        log("إرسال بريد", "resignation", r["id"], f"إلى {', '.join(to_list)} — {subject}")
                        st.success(f"✅ تم الإرسال إلى {len(to_list + cc_list)} مستلم")
                    except smtplib.SMTPAuthenticationError:
                        st.error("❌ فشل تسجيل الدخول. استخدم App Password وتأكد أن SMTP AUTH مفعّل للصندوق.")
                    except Exception as ex:
                        st.error(f"❌ فشل الإرسال: {ex}")

        with right:
            st.markdown('<div class="card"><h4>👁️ معاينة الرسالة</h4>'
                        '<div class="muted">هذا هو الشكل الذي سيصل للمستلم</div></div>',
                        unsafe_allow_html=True)
            prev = f"data:image/png;base64,{img_to_b64(r['photo'])}" if r["photo"] else None
            components.html(build_email_html(r, note, prev, kind), height=860, scrolling=True)

# =========================================================
# TAB: Analytics
# =========================================================
with tab_an:
    if not recs:
        st.info("لا توجد بيانات كافية للتحليل.")
    else:
        df = pd.DataFrame([
            {"branch": r["branch"] or "—", "reason": r["reason"] or "—", "manager": r["direct_manager"] or "—",
             "status": r["status"], "pay": r["net_pay"], "exp": r["exp_days"],
             "month": r["resignation_date"].strftime("%Y-%m")}
            for r in recs
        ])
        a = st.columns(4)
        a[0].markdown(kpi("إجمالي المستحقات", f"{df['pay'].sum():,}", "💰", 22), unsafe_allow_html=True)
        a[1].markdown(kpi("متوسط الخدمة", tenure_text(int(df["exp"].mean())), "⏱️", 18), unsafe_allow_html=True)
        a[2].markdown(kpi("نسبة الاعتماد", f"{(n_appr / n_total * 100):.0f}%", "📈", 22), unsafe_allow_html=True)
        a[3].markdown(kpi("عدد الفروع", df["branch"].nunique(), "🏪", 22), unsafe_allow_html=True)
        st.write("")

        g1, g2 = st.columns(2)
        with g1:
            st.markdown('<div class="card"><h4>الاستقالات حسب الفرع</h4></div>', unsafe_allow_html=True)
            st.bar_chart(df["branch"].value_counts(), color=P["mid"])
        with g2:
            st.markdown('<div class="card"><h4>أسباب الاستقالة</h4></div>', unsafe_allow_html=True)
            st.bar_chart(df["reason"].value_counts(), color=P["accent"])

        g3, g4 = st.columns(2)
        with g3:
            st.markdown('<div class="card"><h4>الاستقالات مقابل التعيينات (شهرياً)</h4></div>', unsafe_allow_html=True)
            res_m = df.groupby("month").size().rename("الاستقالات")
            hire_m = (pd.Series([x["start_date"].strftime("%Y-%m") for x in hires], dtype="object")
                      .value_counts().rename("التعيينات")) if hires else pd.Series(dtype=int, name="التعيينات")
            trend = pd.concat([res_m, hire_m], axis=1).reindex(columns=["الاستقالات", "التعيينات"]).fillna(0).sort_index()
            st.line_chart(trend, color=[P["mid"], P["accent"]])
        with g4:
            st.markdown('<div class="card"><h4>متوسط مدة الخدمة بالأيام (حسب الفرع)</h4></div>', unsafe_allow_html=True)
            st.bar_chart(df.groupby("branch")["exp"].mean().round(0), color=P["brown"])

        g5, g6 = st.columns(2)
        with g5:
            st.markdown('<div class="card"><h4>المستحقات حسب الفرع (EGP)</h4></div>', unsafe_allow_html=True)
            st.bar_chart(df.groupby("branch")["pay"].sum(), color=P["dark"])
        with g6:
            st.markdown('<div class="card"><h4>الاستقالات حسب المدير المباشر</h4></div>', unsafe_allow_html=True)
            st.bar_chart(df["manager"].value_counts(), color=P["mid"])

# =========================================================
# TAB: Audit log
# =========================================================
with tab_log:
    st.markdown('<div class="card"><h4>🧾 سجل النشاط</h4><div class="muted">كل عملية إضافة وتعديل وحذف وإرسال مسجّلة هنا.</div></div>',
                unsafe_allow_html=True)
    logs = fetch("SELECT id, ts, action, entity, entity_id, details FROM audit_log ORDER BY id DESC LIMIT 1000")
    if logs:
        ldf = pd.DataFrame(logs).rename(columns={"id": "#", "ts": "الوقت", "action": "الإجراء",
                                                 "entity": "النوع", "entity_id": "رقم السجل", "details": "التفاصيل"})
        lq = st.text_input("🔎 بحث في السجل", key="log_q")
        if lq:
            mask = ldf.astype(str).apply(lambda col: col.str.contains(lq, case=False, na=False)).any(axis=1)
            ldf = ldf[mask]
        show_df(ldf)
        st.download_button("⬇️ تنزيل السجل CSV", ldf.to_csv(index=False).encode("utf-8-sig"),
                           file_name=f"audit_log_{date.today()}.csv", mime="text/csv", key="dl_log")
    else:
        st.info("لا يوجد نشاط مسجل بعد.")
